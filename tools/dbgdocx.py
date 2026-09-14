import os
from docx import Document

DOCX = r"C:\Users\chuck\Desktop\CA&J Enterprises\agency-agents\caj-growth-system\CA-J-Growth-System-Build-Specs.docx"
doc = Document(DOCX)
print("paragraphs:", len(doc.paragraphs))
print("tables:", len(doc.tables))
seen = {}
for p in doc.paragraphs:
    st = p.style.name if p.style is not None else "<none>"
    seen[st] = seen.get(st, 0) + 1
print("\n--- style histogram ---")
for k, v in sorted(seen.items(), key=lambda x: -x[1])[:14]:
    print("  %-24s %d" % (k, v))

print("\n--- paragraphs whose style contains 'Head' ---")
for p in doc.paragraphs:
    st = p.style.name if p.style is not None else ""
    if "Head" in st:
        print("  [%s] %r" % (st, p.text[:70]))
